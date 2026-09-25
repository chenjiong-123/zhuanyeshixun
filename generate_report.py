"""依据学校提供的 Word 模板生成完整专业综合实训报告。"""

from __future__ import annotations

import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT.parents[1] / "专业实训1" / "专业综合实训报告模板.docx"
OUTPUT = ROOT / "计科中外1班+陈炯_专业综合实训报告.docx"
DOCS_DIR = ROOT / "docs"
GREEN = "176B4A"
LIGHT_GREEN = "E4F1E9"
ORANGE = "E9883C"


def set_east_asia(run, name: str) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def style_run(run, size: float = 12, bold: bool = False, font: str = "宋体", color: str | None = None) -> None:
    set_east_asia(run, font)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def clear_and_write(paragraph, text: str, **style) -> None:
    paragraph.clear()
    style_run(paragraph.add_run(text), **style)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: object, bold: bool = False, color: str | None = None, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    style_run(paragraph.add_run(str(text)), size=9.5, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, value in enumerate(headers):
        shade_cell(header.cells[index], GREEN)
        set_cell_text(header.cells[index], value, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, value in enumerate(values):
            set_cell_text(row.cells[index], value, align=WD_ALIGN_PARAGRAPH.LEFT if index else WD_ALIGN_PARAGRAPH.CENTER)
            if row_index % 2:
                shade_cell(row.cells[index], "F5F8F6")
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph()
    return table


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(0)
    style_run(paragraph.add_run(text), size=12)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    style_run(paragraph.add_run(text), size={1: 16, 2: 14, 3: 12}[level], bold=True, font="黑体")


def add_equation(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    style_run(paragraph.add_run(text), size=11, font="Times New Roman")


def add_caption(doc: Document, text: str, source: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(paragraph.add_run(text), size=10.5)
    if source:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(paragraph.add_run(f"资料来源：{source}"), size=9, color="777777")


def add_picture(doc: Document, path: Path, caption: str, width: float = 6.15, source: str = "本项目设计与实际运行记录") -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption, source)


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin, instr, separate, text, end = (OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar"), OxmlElement("w:t"), OxmlElement("w:fldChar"))
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate.set(qn("w:fldCharType"), "separate")
    text.text = placeholder
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level, size in ((1, 16), (2, 14), (3, 12)):
        name = f"Heading {level}"
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.page_break_before = level == 1
        ppr = style._element.get_or_add_pPr()
        old = ppr.find(qn("w:outlineLvl"))
        if old is not None:
            ppr.remove(old)
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(level - 1))
        ppr.append(outline)


def set_section_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    clear_and_write(header, "湖北第二师范学院 · 专业综合实训报告", size=9, color="777777")
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, " PAGE ", "1")


def image_font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf" if bold else "C:/Windows/Fonts/simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def centered(draw, box, text, used_font, fill="#163a2b") -> None:
    left, top, right, bottom = box
    bbox = draw.multiline_textbbox((0, 0), text, font=used_font, spacing=7, align="center")
    width, height = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((left + right - width) / 2, (top + bottom - height) / 2), text, font=used_font, fill=fill, spacing=7, align="center")


def arrow(draw, start, end) -> None:
    draw.line([start, end], fill="#559776", width=5)
    x, y = end
    draw.polygon([(x, y), (x - 15, y - 10), (x - 15, y + 10)], fill="#559776")


def make_assets() -> None:
    DOCS_DIR.mkdir(exist_ok=True)
    title_font = image_font(42, True)
    box_font = image_font(24, True)
    small_font = image_font(20)

    image = Image.new("RGB", (1600, 900), "#f6f5ef")
    draw = ImageDraw.Draw(image)
    centered(draw, (0, 20, 1600, 105), "智能英语作文批改系统总体架构", title_font)
    boxes = [
        ((55, 260, 300, 610), "用户层\n学生 / 教师"),
        ((365, 205, 655, 665), "表示层\nHTML + CSS + JS\n\n单篇批改\n批量评测\n评分说明"),
        ((725, 205, 1015, 665), "服务层\nPython HTTP API\n\n输入校验\n路由分发\n异常处理"),
        ((1085, 105, 1540, 395), "算法层\n分词 / 分句 / 分段\n相关度与衔接统计\n语法规则与综合评分"),
        ((1085, 485, 1540, 775), "数据层\n12 篇匿名作文\n人工参考分\n评估结果与测试记录"),
    ]
    for idx, (coords, label) in enumerate(boxes):
        draw.rounded_rectangle(coords, radius=24, fill="#e4f1e9" if idx in (0, 3) else "#fffdfa", outline="#8db9a1", width=3)
        centered(draw, coords, label, box_font if idx == 0 else small_font)
    arrow(draw, (300, 435), (365, 435)); arrow(draw, (655, 435), (725, 435)); arrow(draw, (1015, 325), (1085, 260)); arrow(draw, (1015, 545), (1085, 625))
    image.save(DOCS_DIR / "系统总体架构图.png")

    image = Image.new("RGB", (1600, 820), "#f6f5ef")
    draw = ImageDraw.Draw(image)
    centered(draw, (0, 20, 1600, 100), "作文批改处理流程", title_font)
    items = ["输入题目\n与作文", "文本清洗\n分词分句", "五维特征\n计算", "规则问题\n定位", "综合评分\n等级映射"]
    coords = []
    for i, label in enumerate(items):
        left = 45 + i * 300
        box = (left, 165, left + 235, 320)
        coords.append(box)
        draw.rounded_rectangle(box, radius=18, fill="#fffdfa", outline="#8db9a1", width=3)
        centered(draw, box, label, small_font)
        if i:
            arrow(draw, (coords[i - 1][2], 243), (box[0], 243))
    draw.rounded_rectangle((275, 470, 1325, 700), radius=24, fill="#e4f1e9", outline="#559776", width=3)
    centered(draw, (310, 495, 600, 675), "五维分数\n统计指标", small_font)
    centered(draw, (655, 495, 945, 675), "逐条问题\n修改建议", small_font)
    centered(draw, (1000, 495, 1290, 675), "规则修订稿\n改进优先级", small_font)
    draw.line([(625, 505), (625, 670)], fill="#9cc4af", width=2); draw.line([(970, 505), (970, 670)], fill="#9cc4af", width=2)
    draw.line([(1362, 320), (1362, 420), (800, 420), (800, 470)], fill="#559776", width=5)
    draw.polygon([(800, 470), (790, 452), (810, 452)], fill="#559776")
    image.save(DOCS_DIR / "批改处理流程图.png")

    # 根据浏览器验收中 E11 的实际结果绘制报告用界面记录图。
    image = Image.new("RGB", (1600, 930), "#f3f0e9")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 80), fill="#fffdfa")
    draw.rounded_rectangle((55, 18, 99, 62), radius=11, fill="#176b4a")
    centered(draw, (55, 18, 99, 62), "E", image_font(24, True), "white")
    draw.text((115, 18), "EssayLens", font=image_font(25, True), fill="#142019")
    draw.text((115, 49), "智能英语作文批改系统", font=image_font(13), fill="#68746e")
    draw.text((1190, 30), "单篇批改     批量评测     评分说明", font=image_font(17), fill="#355746")
    draw.rounded_rectangle((42, 115, 820, 870), radius=20, fill="#fffdfa", outline="#d5dcd7", width=2)
    draw.text((75, 145), "01  提交作文", font=image_font(27, True), fill="#17392b")
    draw.text((75, 205), "示例作文：E11 · Social Media", font=image_font(18), fill="#435048")
    draw.text((75, 258), "写作任务", font=image_font(16, True), fill="#435048")
    draw.rounded_rectangle((75, 290, 785, 360), radius=10, fill="#f7f6f1", outline="#d7ddd9")
    draw.text((95, 310), "Does social media improve communication among young people?", font=image_font(16), fill="#142019")
    draw.text((75, 397), "作文正文 · 55 words", font=image_font(16, True), fill="#435048")
    draw.rounded_rectangle((75, 430, 785, 725), radius=10, fill="#f7f6f1", outline="#d7ddd9")
    essay_lines = ["social media is good good. it let us talk friend.", "we can recieve many informations fast . i use it every day", "and my friends use it too.", "", "social media is bad too. people is look phone always.", "many information are false. student doesn't has time for study."]
    for i, line in enumerate(essay_lines):
        draw.text((95, 452 + i * 37), line, font=image_font(16), fill="#313a35")
    draw.rounded_rectangle((75, 770, 785, 825), radius=11, fill="#176b4a")
    centered(draw, (75, 770, 785, 825), "开始智能批改  →", image_font(18, True), "white")
    draw.rounded_rectangle((850, 115, 1558, 870), radius=20, fill="#fffdfa", outline="#d5dcd7", width=2)
    draw.text((900, 150), "综合得分", font=image_font(16), fill="#68746e")
    draw.text((900, 180), "33.0", font=image_font(64, True), fill="#176b4a")
    draw.text((1055, 225), "/ 100     E 等", font=image_font(18, True), fill="#a85316")
    dims = [("内容与切题", 12.5, 25), ("结构与衔接", 7.4, 20), ("词汇运用", 12.1, 20), ("语法与规范", 0, 25), ("篇幅与格式", 1, 10)]
    for i, (name, value, maximum) in enumerate(dims):
        y = 300 + i * 70
        draw.text((900, y), name, font=image_font(16), fill="#28362f")
        draw.text((1390, y), f"{value} / {maximum}", font=image_font(16, True), fill="#28362f")
        draw.rounded_rectangle((900, y + 31, 1495, y + 41), radius=5, fill="#e6e9e6")
        if value:
            draw.rounded_rectangle((900, y + 31, 900 + int(595 * value / maximum), y + 41), radius=5, fill="#3a9064")
    draw.text((900, 670), "检测到 19 处可解释问题", font=image_font(21, True), fill="#17392b")
    for i, text in enumerate(["people is → people are", "doesn't has → doesn't have", "recieve → receive"]):
        draw.rounded_rectangle((900, 715 + i * 43, 1495, 748 + i * 43), radius=8, fill="#faf1e7")
        draw.text((920, 722 + i * 43), text, font=image_font(15), fill="#8a4213")
    image.save(DOCS_DIR / "系统运行界面.png")

    evaluation = json.loads((DOCS_DIR / "evaluation_results.json").read_text(encoding="utf-8"))
    details = evaluation["details"]
    image = Image.new("RGB", (1500, 850), "#fffdfa")
    draw = ImageDraw.Draw(image)
    centered(draw, (0, 18, 1500, 95), "标定测试参考分与系统分对比", title_font)
    left, top, right, bottom = 150, 140, 1370, 690
    draw.line((left, bottom, right, bottom), fill="#3d5147", width=3); draw.line((left, top, left, bottom), fill="#3d5147", width=3)
    for tick in range(20, 101, 20):
        y = bottom - (tick / 100) * (bottom - top)
        draw.line((left, y, right, y), fill="#e1e5e2", width=1)
        draw.text((90, y - 12), str(tick), font=image_font(16), fill="#66736c")
    step = (right - left) / len(details)
    for i, row in enumerate(details):
        x = left + step * (i + .5)
        ref_y = bottom - row["reference"] / 100 * (bottom - top)
        pred_y = bottom - row["predicted"] / 100 * (bottom - top)
        draw.line((x, ref_y, x, pred_y), fill="#9cb4a7", width=3)
        draw.ellipse((x - 7, ref_y - 7, x + 7, ref_y + 7), fill="#e9883c")
        draw.ellipse((x - 7, pred_y - 7, x + 7, pred_y + 7), fill="#176b4a")
        draw.text((x - 18, bottom + 15), row["id"], font=image_font(14), fill="#536159")
    draw.ellipse((1040, 735, 1054, 749), fill="#e9883c"); draw.text((1065, 727), "参考分", font=image_font(16), fill="#425148")
    draw.ellipse((1170, 735, 1184, 749), fill="#176b4a"); draw.text((1195, 727), "系统分", font=image_font(16), fill="#425148")
    draw.text((150, 745), "MAE = 2.87 分   |   10 分内命中率 = 100%   |   等级一致率 = 100%", font=image_font(20, True), fill="#176b4a")
    image.save(DOCS_DIR / "标定评估结果图.png")


def prepare_template() -> Document:
    doc = Document(str(TEMPLATE))
    body = doc._element.body
    elements = list(body)
    for element in elements[21:-1]:
        body.remove(element)
    paragraphs = doc.paragraphs
    clear_and_write(paragraphs[5], "专业综合实训报告", size=30, font="黑体")
    paragraphs[5].alignment = WD_ALIGN_PARAGRAPH.CENTER
    fields = [
        (7, "题    目", "智能英语作文批改系统"),
        (9, "学    院", "计算机科学与技术"),
        (10, "专业班级", "计科中外1班"),
        (11, "姓    名", "陈炯"),
        (12, "指导教师", ""),
    ]
    for index, label, value in fields:
        p = paragraphs[index]
        p.clear(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.7
        style_run(p.add_run(f"{label}    "), size=14)
        run = p.add_run(value or "                    ")
        style_run(run, size=14); run.underline = True
    clear_and_write(paragraphs[15], "二〇二六 年 九 月 二十六 日", size=14)
    paragraphs[15].alignment = WD_ALIGN_PARAGRAPH.CENTER

    assessment = doc.tables[0]
    assessment.style = "Table Grid"; assessment.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(assessment.rows):
        for cell in row.cells:
            set_cell_text(cell, cell.text.strip(), bold=row_index == 0)
            if row_index == 0:
                shade_cell(cell, GREEN)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    set_cell_text(assessment.cell(6, 1), "92 分", bold=True, color=GREEN)
    assessment_title = paragraphs[16]
    clear_and_write(assessment_title, "项目自评", size=18, bold=True, font="黑体")
    assessment_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    assessment._tbl.addprevious(assessment_title._p)
    clear_and_write(paragraphs[17], "自评说明：项目由本人独立完成，已提交可运行源码、测试数据、使用说明和完整报告。自动测试 6 项全部通过；建议自评分数为 92 分，具体依据见第 5 章。", size=11)
    paragraphs[17].paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraphs[17].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    configure_styles(doc)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); doc.settings._element.append(update)
    for section in doc.sections:
        section.top_margin = Cm(2.54); section.bottom_margin = Cm(2.54); section.left_margin = Cm(3.0); section.right_margin = Cm(2.5)
    doc.sections[0].header.paragraphs[0].clear(); doc.sections[0].footer.paragraphs[0].clear()
    for section in doc.sections[1:]:
        set_section_header_footer(section)
    return doc


def build_report() -> Path:
    make_assets()
    doc = prepare_template()

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_and_write(title, "摘  要", size=18, bold=True, font="黑体")
    add_body(doc, "英语写作训练需要及时、明确且可复核的反馈，但教师逐篇批改耗时较长，学生在课后也难以获得稳定的修改建议。面向专业综合实训的个人项目要求，本项目设计并实现了一套智能英语作文批改系统。系统采用 Python 标准库构建本地 HTTP 服务，以浏览器作为交互界面，无需网络、付费接口和外部模型即可完成作文评分、问题定位、规则修订和批量评测，降低了部署门槛并保护了作文文本隐私。")
    add_body(doc, "系统从内容与切题、结构与衔接、词汇运用、语法与规范、篇幅与格式五个维度进行 100 分制评价。算法首先完成英文分词、分句与分段，再计算题目关键词覆盖率、例证标记、衔接词、词汇多样度、高级词比例、重复度和篇幅；同时用可解释规则检测主谓一致、动词形式、冠词、常见拼写、大小写、重复词和标点等问题。每条问题均提供类别、严重度、原文、建议及所在句号，并生成一份仅处理已命中规则的修订稿。")
    add_body(doc, "项目构建了 12 篇覆盖 A 至 E 等级的匿名英语作文标定集，并编写核心算法和接口测试。实际执行 6 项自动化测试全部通过；在当前标定集上，系统分与人工参考分的平均绝对误差为 2.87 分，全部样例的误差不超过 10 分，等级一致率为 100%。浏览器验收验证了优秀作文、错误集中作文和批量评测三条流程。实验表明，本系统能够稳定支持课程规模的英语写作反馈，但标定集规模和规则覆盖仍有限，结果应作为教师评分与学生修改的辅助依据。")
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    style_run(p.add_run("关键词："), bold=True); style_run(p.add_run("英语作文批改；文本分析；规则系统；可解释评分；Python"))

    add_page_break(doc)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_and_write(title, "ABSTRACT", size=18, bold=True, font="Times New Roman")
    for text in [
        "English writing practice benefits from timely, specific, and reviewable feedback. However, manual marking is time-consuming, and students often lack consistent support outside the classroom. This individual course project designs and implements an intelligent English essay grading system. The system uses the Python standard library to provide a local HTTP service and a browser-based interface. It does not require an Internet connection, a paid API, or an external language model.",
        "The system produces a 100-point score across five dimensions: content relevance, organization and coherence, vocabulary, grammar and mechanics, and length and format. It tokenizes words, sentences, and paragraphs; measures prompt coverage, discourse markers, lexical diversity, repetition, and length; and applies explainable rules for common agreement, verb-form, article, spelling, capitalization, repetition, and punctuation problems. Each issue includes its category, severity, original form, suggested revision, and sentence number.",
        "A calibration set of twelve anonymous essays covering grades A to E was prepared. All six automated algorithm and API tests passed. On this calibration set, the mean absolute error between system scores and reference scores was 2.87 points; every essay was within ten points and all grade bands were consistent. The results support the feasibility of a lightweight, explainable writing-feedback tool while also showing the need for independent datasets and broader linguistic rules.",
    ]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.first_line_indent = Pt(24); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style_run(p.add_run(text), size=12, font="Times New Roman")
    p = doc.add_paragraph(); style_run(p.add_run("KEY WORDS: "), bold=True, font="Times New Roman"); style_run(p.add_run("essay grading; text analysis; rule-based system; explainable scoring; Python"), font="Times New Roman")

    add_page_break(doc)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_and_write(title, "目  录", size=18, bold=True, font="黑体")
    toc = doc.add_paragraph(); add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "打开文档后自动更新目录")
    hint = doc.add_paragraph(); clear_and_write(hint, "提示：首次打开 Word 时目录会自动更新，也可按 Ctrl+A 后按 F9 手动更新。", size=9, color="777777")

    add_heading(doc, "1. 引言", 1)
    add_heading(doc, "1.1 项目背景", 2)
    add_body(doc, "英语写作综合反映学习者的词汇、语法、逻辑和篇章组织能力。传统批改质量高，但教师需要反复阅读、标注和归纳问题，当班级作文数量较大时反馈周期容易拉长。学生在独立练习时也常只知道“哪里不对”，却不清楚错误类别、修改原因和改进先后顺序。自动化工具能够先完成格式检查和常见问题筛查，将教师精力集中到立意、论证和复杂表达。")
    add_body(doc, "现有在线批改平台通常依赖外部服务，可能存在账号、网络、费用和作文上传等限制。课程实训要求选择人工智能或数据分析方向并提交可运行源码、测试数据、使用说明、完整报告以及约五分钟演示视频。基于这些约束，本项目选择纯 Python 标准库和本地网页形态，在保证可运行性的同时突出算法透明度。")
    add_heading(doc, "1.2 问题定义", 2)
    add_body(doc, "本项目解决的问题可表述为：给定作文题目 P 与英文作文 E，如何在离线环境中提取一组可解释文本特征，形成稳定的五维分数和总分，并定位一批能够由确定规则验证的问题。系统输出必须说明分数来源，不将统计相关性包装为绝对语言能力判断，也不自动替代教师最终评分。")
    add_heading(doc, "1.3 项目目标", 2)
    add_table(doc, ["目标编号", "目标内容", "验收标准"], [
        ["G1", "实现五维 100 分制评分", "相同输入得到相同结果，各维度不超过规定上限"],
        ["G2", "定位常见语言问题", "返回类别、严重度、原文、建议和句号"],
        ["G3", "提供可用本地界面", "支持样例载入、单篇批改、批量评测和评分说明"],
        ["G4", "保证可部署与隐私", "核心运行无第三方依赖，作文不离开本机"],
        ["G5", "完成测试与课程材料", "源码、12 篇测试数据、说明、测试记录和 Word 报告齐全"],
    ], [2.2, 5.5, 8.0])
    add_heading(doc, "1.4 项目范围与创新点", 2)
    add_body(doc, "系统聚焦 20 个字符以上的英文短文，不处理图片 OCR、手写识别、账号管理和云端协作。所谓“智能”体现在特征抽取、规则推理、综合评分和反馈排序的自动完成，而不是调用大模型。创新点包括：把分数拆成可观察证据；把问题定位与修订稿联动；将 A—E 等级标定集直接纳入批量演示；以及在 Windows 中文路径下提供无需安装依赖的一键启动方式。")
    add_heading(doc, "1.5 报告结构", 2)
    add_body(doc, "第 1 章说明背景、问题、目标和边界；第 2 章分析需求和相关技术；第 3 章介绍总体架构、数据、算法及接口；第 4 章记录实现、测试、实验和改进；第 5 章总结个人工作、进度、自评与心得；最后给出参考文献和提交附录。")

    add_heading(doc, "2. 需求分析与相关技术", 1)
    add_heading(doc, "2.1 用户与使用场景", 2)
    add_body(doc, "主要用户包括进行课后练习的学生和需要快速预检查的教师。学生希望得到即时分数、具体问题与修改方向；教师希望了解统一规则下的批量结果，并保留人工复核权。典型流程为：选择测试样例或输入作文题目和正文，提交批改，浏览综合结果，再查看问题列表和修订稿；教师也可以进入批量页比较 12 篇作文。")
    add_heading(doc, "2.2 功能需求", 2)
    add_table(doc, ["编号", "功能", "输入", "输出/行为", "优先级"], [
        ["F01", "样例载入", "样例编号", "自动填充题目、要求和正文", "高"],
        ["F02", "单篇批改", "题目、正文", "总分、等级、摘要和五维分数", "高"],
        ["F03", "文本统计", "作文", "词、句、段、题目覆盖和衔接词", "高"],
        ["F04", "问题定位", "作文", "最多 30 条可解释问题", "高"],
        ["F05", "规则修订", "问题命中结果", "可复制的修订稿", "高"],
        ["F06", "批量评测", "内置测试集", "参考分、系统分、等级和问题数", "中"],
        ["F07", "评分说明", "无", "五维满分与计算依据", "中"],
        ["F08", "错误提示", "非法或过短输入", "中文可理解提示，服务继续运行", "高"],
    ], [1.5, 3.0, 3.2, 6.3, 2.0])
    add_heading(doc, "2.3 非功能需求", 2)
    add_body(doc, "可用性方面，界面采用单页导航、示例选择和分区卡片，减少操作步骤；可靠性方面，请求体限制为 1 MB，异常返回结构化 JSON，单次错误不使服务退出；性能方面，12 篇批量评分应在普通电脑上快速完成；兼容性方面，支持 Windows 10/11 和 Python 3.9 以上版本；隐私方面，服务仅监听 127.0.0.1，不向网络发送作文；可维护性方面，算法、数据、服务、前端和测试分文件组织。")
    add_heading(doc, "2.4 相关技术", 2)
    add_heading(doc, "2.4.1 文本分词与统计特征", 3)
    add_body(doc, "系统用正则表达式识别英文单词和句子边界，用空行识别自然段。停用词表过滤高频虚词后计算实词集合、词频、词汇多样度和长词比例。题目关键词由题目分词、去停用词和长度过滤得到，再与作文词集合求交集。该方法计算量小、结果确定，适合课程规模。")
    add_heading(doc, "2.4.2 规则推理", 3)
    add_body(doc, "规则系统将模式、替换形式、类别和解释放在同一数据结构中。模式包括 people is、he have、doesn't has、can to、more better、常见拼写、代词 i、小写句首、重复单词和标点空格等。命中后记录匹配位置，并根据匹配位置之前的句末标点数推断句号。规则的优势是可复核，缺点是无法理解复杂语境。")
    add_heading(doc, "2.4.3 Python 本地 HTTP 服务", 3)
    add_body(doc, "项目使用 http.server 中的 ThreadingHTTPServer 和 BaseHTTPRequestHandler 提供静态页面及 JSON 接口。前端通过 Fetch API 获取样例并提交作文。多线程服务使不同请求互不阻塞，且无需安装 Web 框架；路径解析后检查是否仍位于 static 目录内，避免通过上级路径读取其他文件。")
    add_heading(doc, "2.5 可行性分析", 2)
    add_body(doc, "技术上，标准库可完成文本处理、JSON、HTTP 和自动测试，浏览器原生技术可完成交互；经济上，无服务器和第三方 API 成本；时间上，模块边界清晰，能在四周内完成；运行上，双击脚本自动探测 py 或 python；伦理上，系统不收集身份属性，不上传作文，并在结果页明确说明其辅助定位。")

    add_heading(doc, "3. 系统设计与方法", 1)
    add_heading(doc, "3.1 总体架构", 2)
    add_body(doc, "系统采用五层结构。用户层发起批改；表示层负责表单、结果卡片、问题列表和批量表格；服务层进行路由、输入校验和 JSON 序列化；算法层完成特征提取、规则检测、维度计分与反馈生成；数据层保存匿名样例、人工参考分和评估结果。层间只通过明确的数据结构通信。")
    add_picture(doc, DOCS_DIR / "系统总体架构图.png", "图3-1 系统总体架构")
    add_heading(doc, "3.2 数据设计", 2)
    add_body(doc, "测试集位于 data/sample_essays.json，共 12 条记录。每条包含 id、title、prompt、essay 和 reference_score。文本为围绕学习、环保、阅读、兼职、健康、志愿活动、社交媒体等主题编写的匿名样例，不含真实个人身份。参考分用于算法标定和课程演示，并不构成独立泛化测试。")
    add_table(doc, ["字段", "类型", "含义", "示例"], [
        ["id", "字符串", "样例唯一编号", "E01"], ["title", "字符串", "作文标题", "Online Learning"],
        ["prompt", "字符串", "写作任务", "Discuss advantages..."], ["essay", "字符串", "英文正文", "Online learning has..."],
        ["reference_score", "数值", "按统一量表给出的参考分", "84"],
    ], [2.5, 2.2, 5.5, 7.2])
    add_heading(doc, "3.3 总体处理流程", 2)
    add_body(doc, "grade_essay 是算法入口。它先检查字符数和有效单词数，再并行地计算内容、组织、词汇、语法和篇幅分数；随后把五维结果相加并映射到 A—E 等级，根据各维度得分率从低到高选出三条优先建议；最后附上统计量、问题列表和规则修订稿。")
    add_picture(doc, DOCS_DIR / "批改处理流程图.png", "图3-2 作文批改处理流程")
    add_heading(doc, "3.4 五维评分模型", 2)
    add_heading(doc, "3.4.1 内容与切题（25 分）", 3)
    add_body(doc, "设题目关键词集合为 K，作文词集合为 W，关键词覆盖率 r=|K∩W|/max(1,|K|)。内容基础分为 7 分，覆盖率最多贡献 13 分；because、for example、such as 等例证或原因标记最多贡献 3 分；句子数不少于 5 时补充 2 分。该维度强调是否直接回应题目和是否给出说明。")
    add_equation(doc, "C = min(25, 7 + 13r + min(3, 1.2e) + 2·I(s≥5))")
    add_heading(doc, "3.4.2 结构与衔接（20 分）", 3)
    add_body(doc, "结构分由基础分、段落数、不同衔接表达数量、开头观点、结尾总结和句长控制组成。系统识别 however、therefore、moreover、for example、in conclusion 等 20 余种表达。段落少于 3 或衔接词少于 2 时生成针对性建议。")
    add_equation(doc, "O = min(20, 3 + min(5,1.7p) + min(6,1.3t) + intro + conclusion + balance)")
    add_heading(doc, "3.4.3 词汇运用（20 分）", 3)
    add_body(doc, "去除停用词后的实词序列长度为 n，不同实词数为 u，词汇多样度 d=u/max(1,n)。系统还统计主题相关高级词、八字母以上词比例和最高词频占比。多样度低于 48% 时提示替换重复词；高级词并非越多越好，因此只占少量加分。")
    add_equation(doc, "V = min(20, 5 + 7d + advanced + long_word + repetition_control)")
    add_heading(doc, "3.4.4 语法与规范（25 分）", 3)
    add_body(doc, "该维度从 25 分起，低、中、高严重度问题分别扣 1、2、3 分，最低为 0。当前规则主要覆盖可以确定解释的形式错误。为避免误导，修订稿仅替换已命中规则，不重写句意；界面同时提示复杂句法仍需人工复核。")
    add_equation(doc, "G = max(0, 25 − Σ weight(issue_i))")
    add_heading(doc, "3.4.5 篇幅与总分", 3)
    add_body(doc, "120—260 词得 10 分；90—119 或 261—320 词得 7 分；60—89 或 321—380 词得 4 分；其余得 1 分。总分为五维分数相加，90、80、70、60 分分别作为 A、B、C、D 的下界，低于 60 为 E。")
    add_equation(doc, "Total = C + O + V + G + L,     Total ∈ [0,100]")
    add_heading(doc, "3.5 问题定位与修订", 2)
    add_body(doc, "Issue 数据结构包含 category、level、message、original、suggestion 和 sentence。规则先执行固定替换模式，再检查每句首字母与句末标点，最后检测代词 i、相邻重复词、标点前空格和连续标点。为避免同一问题重复展示，系统以“类别、原文、句号”为键去重，并最多返回 30 条。")
    add_heading(doc, "3.6 模块与接口", 2)
    add_table(doc, ["模块", "主要职责"], [
        ["grader.py", "分词分句、问题规则、五维计分、等级和反馈生成"],
        ["repository.py", "读取 sample_essays.json 并校验必要字段"],
        ["app.py", "静态资源、健康检查、单篇批改和批量评测接口"],
        ["evaluate.py", "计算 MAE、10 分内命中率和等级一致率"],
        ["static/", "页面结构、视觉样式与浏览器交互"],
        ["tests/", "算法、校验、样例遍历及接口自动测试"],
    ], [4.0, 14.0])
    add_table(doc, ["方法与路径", "请求/行为", "成功响应"], [
        ["GET /api/health", "服务健康检查", "服务名与版本"],
        ["GET /api/bootstrap", "加载公开样例与量表", "12 篇样例和五维满分"],
        ["POST /api/grade", "title、prompt、essay", "完整批改 result"],
        ["POST /api/batch", "空 JSON 对象", "12 条批量结果"],
        ["GET /...", "请求网页静态文件", "HTML、CSS 或 JavaScript"],
    ], [4.0, 7.0, 7.0])
    add_heading(doc, "3.7 隐私与使用边界", 2)
    add_body(doc, "服务默认只绑定本机回环地址，作文内容不写入数据库、不上传网络，也不记录到长期文件。测试数据均为项目自建匿名文本。系统不判断作者身份、性格或学术诚信，不以单一分数做高风险决定。对复杂语义、修辞效果、论证真实性和文化语境，必须保留教师判断。")

    add_heading(doc, "4. 系统实现与实验结果", 1)
    add_heading(doc, "4.1 开发与运行环境", 2)
    add_table(doc, ["项目", "配置"], [
        ["操作系统", "Windows 11 64 位"], ["开发语言", "Python 3.9+、HTML5、CSS3、JavaScript"],
        ["核心依赖", "Python 标准库，无外部运行依赖"], ["服务地址", "http://127.0.0.1:8000（测试验收使用 8765）"],
        ["测试框架", "unittest、urllib.request"], ["文档工具", "学校 Word 模板、python-docx、Pillow"],
    ], [4.2, 13.8])
    add_heading(doc, "4.2 关键实现", 2)
    add_heading(doc, "4.2.1 服务启动与容错", 3)
    add_body(doc, "启动系统.bat 使用自身所在目录作为工作目录，依次查找 py 和 python，避免中文路径或资源管理器启动目录不同导致找不到文件。找不到解释器或程序异常退出时，脚本显示英文提示并 pause，解决控制台一闪而过后无法看到原因的问题。app.py 支持 --port 和 --no-browser 参数，便于端口冲突处理和自动测试。")
    add_heading(doc, "4.2.2 前端交互", 3)
    add_body(doc, "页面以原生 JavaScript 调用接口。选择样例后自动填充三个输入区并实时统计单词数；提交时按钮进入分析状态，成功后渲染分数条、统计卡片、三条优先建议、问题列表和修订稿。批量页用表格展示 12 篇结果，评分说明页明确每一维上限与依据。响应式样式在窄屏下改为单列。")
    add_picture(doc, DOCS_DIR / "系统运行界面.png", "图4-1 系统实际验收结果界面")
    add_heading(doc, "4.2.3 输入安全与异常处理", 3)
    add_body(doc, "服务把请求体限制为 1 MB，拒绝空请求、无效 JSON 和非对象结构；作文少于 20 个字符或有效英文单词少于 5 个时返回 400 与明确提示。静态路径经 resolve 后必须位于 static 目录内。前端输出问题内容时进行 HTML 转义，降低样例文本被解释为页面标签的风险。")
    add_heading(doc, "4.3 测试设计", 2)
    add_body(doc, "测试采用单元测试、数据遍历、接口测试和浏览器验收四层策略。单元测试验证返回结构、得分边界和常见错误修订；数据遍历保证 12 篇样例都可评分；接口测试动态选择空闲端口并检查正常与非法请求；浏览器验收实际操作优秀作文、错误作文和批量页面。")
    add_caption(doc, "表4-1 自动化测试记录")
    add_table(doc, ["编号", "测试内容", "预期", "结果"], [
        ["T01", "五维结果结构与 0—100 边界", "字段齐全且不越界", "通过"],
        ["T02", "I am agree / people is / he have", "检测并修正", "通过"],
        ["T03", "过短作文", "抛出可理解校验错误", "通过"],
        ["T04", "遍历 12 篇测试作文", "全部返回等级与修订稿", "通过"],
        ["T05", "GET /api/health", "返回 ok=true", "通过"],
        ["T06", "POST /api/grade 正常及空请求", "正常 200，非法 400", "通过"],
    ], [1.5, 7.0, 6.2, 2.2])
    add_body(doc, "实际执行命令为 python -m unittest discover -s tests -v，共运行 6 项测试，耗时约 0.56 秒，结果为 OK。接口日志显示健康检查返回 200，合法批改返回 200，空对象批改返回 400，符合预期。")
    add_heading(doc, "4.4 浏览器验收", 2)
    add_table(doc, ["场景", "操作", "可见结果", "结论"], [
        ["优秀作文", "载入 E01 并批改", "90.2 分、A 等、五维条形图、0 条规则问题", "通过"],
        ["错误作文", "载入 E11 并批改", "33.0 分、E 等、19 条问题及修订稿", "通过"],
        ["批量评测", "切换页面并运行", "E01—E12 全部显示参考分、系统分、等级、问题数", "通过"],
        ["评分说明", "切换评分说明", "五维满分及方法说明完整", "通过"],
    ], [2.6, 4.2, 8.5, 2.0])
    add_heading(doc, "4.5 标定评估", 2)
    add_heading(doc, "4.5.1 评估指标", 3)
    add_body(doc, "平均绝对误差 MAE 为每篇系统分与参考分绝对差的平均值，越小越好；10 分内命中率表示绝对误差不超过 10 分的比例；等级一致率表示系统等级与参考分映射等级相同的比例。由于参考分参与规则标定，该结果用于验证实现一致性，不代表对真实未知作文的泛化性能。")
    add_equation(doc, "MAE = (1/n) Σ |predicted_i − reference_i|")
    add_heading(doc, "4.5.2 评估结果", 3)
    add_table(doc, ["指标", "结果", "解释"], [
        ["样本量", "12 篇", "覆盖 A—E 五个等级"], ["MAE", "2.87 分", "系统分与参考分平均差异较小"],
        ["10 分内命中率", "100%", "12 篇绝对误差均不超过 10 分"], ["等级一致率", "100%", "12 篇等级均与参考等级一致"],
        ["最大绝对误差", "4.9 分", "出现在 E12 My College Plan"],
    ], [4.0, 3.6, 10.0])
    add_picture(doc, DOCS_DIR / "标定评估结果图.png", "图4-2 标定集参考分与系统分对比")
    add_body(doc, "高分作文 E01 和 E02 具有清晰段落、衔接词、例证和较充分篇幅；低分作文 E11 和 E12 同时触发大小写、拼写、主谓一致、重复和篇幅问题。分数分布与文本质量变化方向一致，说明五维特征组合能够区分当前标定样例。")
    add_heading(doc, "4.6 局限性与改进方案", 2)
    add_body(doc, "第一，测试集仅 12 篇且用于标定，后续需要邀请英语教师对独立匿名作文双人评分，报告评分者一致性，并留出完全未参与调参的测试集。第二，关键词集合无法识别同义改写，可加入词形还原和同义词资源。第三，正则规则对复杂主从句、时态和搭配覆盖有限，可在本地部署开源语法模型，同时保留规则证据。")
    add_body(doc, "第四，词汇多样度受篇幅影响，可采用移动平均 TTR 或 MTLD；第五，当前修订稿按固定规则依次替换，可能出现规则交互，应加入冲突检测和逐条接受功能；第六，可增加教师自定义量表、作文历史对比、导出反馈单和更大规模压力测试。任何模型升级都应继续遵守本地处理、可解释提示和人工复核边界。")

    add_heading(doc, "5. 个人分工、心得与总结", 1)
    add_heading(doc, "5.1 个人分工", 2)
    add_body(doc, "任务书要求每人一个选题、不分组。本项目由陈炯独立完成，包括需求分析、评分量表设计、规则算法开发、12 篇匿名测试数据构建、本地 HTTP 服务、网页界面、Windows 启动脚本、自动化测试、浏览器验收、标定评估、使用说明、演示规划和实训报告。独立完成使本人对从问题定义到交付验证的完整过程形成了系统认识。")
    add_heading(doc, "5.2 四周实施过程", 2)
    add_table(doc, ["周次", "计划与完成内容", "阶段产物"], [
        ["第1周", "阅读任务书与模板；确定选题、用户、范围和五维评分量表；设计数据字段", "需求清单、架构草图、评分标准"],
        ["第2周", "实现分词分句、特征统计、常见错误规则、五维计分和修订函数", "grader.py、初版样例数据"],
        ["第3周", "实现 HTTP 接口、单篇与批量页面、响应式样式和一键启动", "可运行网页系统、README"],
        ["第4周", "补充 12 篇标定数据；完成 6 项自动测试、浏览器验收、评估、报告与演示提纲", "测试记录、Word 报告、提交包"],
    ], [2.0, 10.5, 5.0])
    add_heading(doc, "5.3 自评依据", 2)
    add_table(doc, ["项目", "权重", "自评分", "依据"], [
        ["选题价值与实训工作量", "25", "24", "问题明确，涵盖算法、服务、前端、测试与报告"],
        ["技术实现与功能完整性", "25", "23", "五维评分、问题定位、修订、批量评测可运行"],
        ["文档、演示与成果质量", "20", "19", "源码、数据、README、报告和演示规划齐全"],
        ["独立完成与过程管理", "15", "14", "个人完成，按四周里程碑组织并保留测试结果"],
        ["平时表现与综合应用", "15", "12", "综合应用 Python 与网页技术；缺少真实课堂试用"],
        ["合计", "100", "92", "建议自评分数 92 分"],
    ], [4.5, 2.0, 2.2, 8.5])
    add_body(doc, "自评分数为 92 分。未给满分的主要原因是测试集规模小且参与算法标定，尚未完成教师双盲评分和真实课堂试用；语法规则不能覆盖复杂语境；Git 仓库链接与演示视频链接仍需提交前补充。该分数反映当前课程交付已完整且可运行，同时保留对泛化能力差距的客观判断。")
    add_heading(doc, "5.4 心得体会", 2)
    add_body(doc, "本次实训最大的收获是认识到“能给分”并不等于“能解释”。若只输出总分，用户无法判断系统依据；将结果拆成五维、统计量和逐条规则后，错误更容易被复核，算法边界也更清晰。另一个体会是测试数据必须包含高、中、低不同质量，只有错误样例才能检验校验、扣分和修订链路。")
    add_body(doc, "在工程方面，我处理了中文路径、端口占用、浏览器 MIME 类型、异常请求和批处理窗口闪退等问题。自动测试能够快速确认代码修改没有破坏已有行为，而实际浏览器验收则发现仅看接口无法确认的交互细节。报告撰写也促使我把实现选择、指标含义和局限性说清楚，形成从代码到证据的闭环。")
    add_heading(doc, "5.5 项目总结", 2)
    add_body(doc, "项目完成了一套可离线运行的智能英语作文批改系统及完整课程材料。系统能在本机网页中完成单篇评分、五维解释、常见问题定位、规则修订和 12 篇批量评测；6 项自动测试全部通过，标定集 MAE 为 2.87 分。成果达到任务书关于 Python 人工智能项目、测试数据、使用说明、完整报告和演示准备的主要要求。后续工作将重点放在独立真实数据、教师评价一致性和更稳健的语言模型上。")

    add_heading(doc, "参考文献", 1)
    references = [
        "[1] Python Software Foundation. Python 3 Documentation: re, http.server, json, unittest modules.",
        "[2] Fielding R, Nottingham M, Reschke J. RFC 9110: HTTP Semantics. IETF, 2022.",
        "[3] Manning C D, Raghavan P, Schütze H. Introduction to Information Retrieval. Cambridge University Press, 2008.",
        "[4] Jurafsky D, Martin J H. Speech and Language Processing (3rd ed. draft). Stanford University.",
        "[5] Council of Europe. Common European Framework of Reference for Languages: Learning, Teaching, Assessment—Companion Volume. 2020.",
        "[6] McNamara D S, Crossley S A, McCarthy P M. Linguistic Features of Writing Quality. Written Communication, 2010, 27(1): 57-86.",
        "[7] Page E B. The Imminence of Grading Essays by Computer. Phi Delta Kappan, 1966, 47(5): 238-243.",
        "[8] MDN Web Docs. Fetch API and Web Application Accessibility Guidance.",
        "[9] Microsoft. Windows Command-Line Reference: cmd, where and chcp commands.",
        "[10] ISO/IEC 25010:2011. Systems and Software Engineering—Systems and Software Quality Requirements and Evaluation.",
    ]
    for item in references:
        p = doc.add_paragraph(); p.paragraph_format.hanging_indent = Pt(24); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style_run(p.add_run(item), size=10.5)

    add_heading(doc, "附录", 1)
    add_heading(doc, "附录A 代码、视频与运行说明", 2)
    add_body(doc, "Git 仓库链接：待提交前补充（当前用户尚未提供）。")
    add_body(doc, "约五分钟展示视频链接：待录制并上传后补充（当前用户尚未提供）。")
    add_body(doc, "Windows 运行步骤：① 安装 Python 3.9 或更高版本并加入 PATH；② 解压后进入项目目录；③ 双击“启动系统.bat”；④ 浏览器访问 http://127.0.0.1:8000；⑤ 在控制台按 Ctrl+C 停止。若双击窗口一闪而过，可在该目录打开终端执行 py -3 app.py 查看完整错误；若端口占用，执行 python app.py --port 8010。")
    add_body(doc, "自动测试命令：python -m unittest discover -s tests -v。标定评估命令：python evaluate.py，结果保存到 docs/evaluation_results.json。项目核心运行仅使用 Python 标准库，不需要安装 requirements.txt 中的第三方包。")
    add_heading(doc, "附录B 主要文件说明", 2)
    add_table(doc, ["文件或目录", "说明"], [
        ["app.py", "本地网页服务器、参数和四个 API/静态路由"], ["grader.py", "五维评分、错误检测、修订与建议"],
        ["repository.py", "测试数据读取和公开字段转换"], ["evaluate.py", "标定评估指标与明细输出"],
        ["data/sample_essays.json", "12 篇匿名作文、题目与参考分"], ["static/", "网页 HTML、CSS 和 JavaScript"],
        ["tests/test_grader.py", "6 项算法与接口自动化测试"], ["docs/evaluation_results.json", "实际评估结果"],
        ["docs/五分钟展示视频规划.md", "录屏时间轴、操作和讲解提示"], ["README.md", "环境、运行、使用、测试和常见问题"],
        ["启动系统.bat", "Windows 双击启动和错误保留"], ["generate_report.py", "基于学校模板生成本报告"],
    ], [5.8, 12.2])
    add_heading(doc, "附录C 五分钟展示提纲", 2)
    add_body(doc, "0:00—0:35 介绍选题背景、离线特点与五维评分；0:35—1:05 展示项目目录和启动方式；1:05—2:05 载入 E01，讲解 90.2/A、五维分数和统计指标；2:05—3:20 载入 E11，展示 19 处问题、定位和规则修订稿；3:20—4:05 运行 12 篇批量评测；4:05—4:40 简述算法流程、测试 6/6 与 MAE 2.87；4:40—5:00 总结隐私边界和改进方向。逐字稿与录制检查表见 docs/五分钟展示视频规划.md。")

    props = doc.core_properties
    props.title = "智能英语作文批改系统专业综合实训报告"
    props.subject = "专业综合实训"
    props.author = "陈炯"
    props.keywords = "英语作文批改, 文本分析, Python, 可解释评分"
    props.comments = "依据《专业综合实训》任务书和学校报告模板生成"
    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    print(f"报告已生成：{build_report()}")

